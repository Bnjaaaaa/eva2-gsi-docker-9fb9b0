"""Genera el informe de la Evaluación 2 (Word) con formato APA 6.

Uso: python generar_informe.py <salida.docx>
Las leyendas usan campos SEQ y los índices campos TOC: Word los completa al ejecutar word_actualizar.ps1.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
CAP = RAIZ / "capturas"
DIA = RAIZ / "diagramas"
salida = Path(sys.argv[1])

FUENTE = "Times New Roman"
ROJO = RGBColor(0xC0, 0x00, 0x00)
COLOR_RIESGO = {"Alto": ("C00000", "FFFFFF"), "Medio": ("FFC000", "000000"), "Bajo": ("92D050", "000000"),
                "Informativo": ("5B9BD5", "FFFFFF")}

doc = Document()

# ---------------------------------------------------------------- estilos y página
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)          # tamaño carta
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(2.54))


def fijar_fuente(estilo, tam, negrita=None, cursiva=None, color=RGBColor(0, 0, 0)):
    estilo.font.name = FUENTE
    estilo.font.size = Pt(tam)
    estilo.font.color.rgb = color
    if negrita is not None:
        estilo.font.bold = negrita
    if cursiva is not None:
        estilo.font.italic = cursiva
    rpr = estilo.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FUENTE)


normal = doc.styles["Normal"]
fijar_fuente(normal, 12)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 2.0

for nombre, centrado, cursiva in (("Heading 1", True, False), ("Heading 2", False, False), ("Heading 3", False, True)):
    e = doc.styles[nombre]
    fijar_fuente(e, 12, negrita=True, cursiva=cursiva)
    pf = e.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before, pf.space_after, pf.line_spacing = Pt(12), Pt(6), 1.5
    pf.keep_with_next = True

cap = doc.styles["Caption"]
fijar_fuente(cap, 11, negrita=False, cursiva=False)
cap.paragraph_format.line_spacing = 1.0
cap.paragraph_format.space_after = Pt(12)

from docx.enum.style import WD_STYLE_TYPE
for nombre in ("toc 1", "toc 2", "toc 3", "table of figures"):     # estilos que Word usa al generar los índices
    e = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
    e.element.set(qn("w:styleId"), {"toc 1": "TOC1", "toc 2": "TOC2", "toc 3": "TOC3",
                                    "table of figures": "TableofFigures"}[nombre])   # Word los busca por este id
    e.base_style = normal
    fijar_fuente(e, 12)
    e.paragraph_format.line_spacing = 1.15
    e.paragraph_format.space_after = Pt(4)
    if nombre in ("toc 2", "toc 3"):
        e.paragraph_format.left_indent = Cm(0.6 if nombre == "toc 2" else 1.2)

for nombre in ("List Bullet", "List Number"):
    fijar_fuente(doc.styles[nombre], 12)
    doc.styles[nombre].paragraph_format.line_spacing = 2.0

# número de página arriba a la derecha (APA 6)
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def campo(parrafo, instr, texto="1", fmt=None):
    r = parrafo.add_run()
    if fmt:
        fmt(r)
    for tipo, val in (("begin", None), ("instr", instr), ("separate", None), ("text", texto), ("end", None)):
        if tipo in ("begin", "separate", "end"):
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        elif tipo == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = f" {val} "
        else:
            e = OxmlElement("w:t")
            e.text = val
        r._r.append(e)
    return r


campo(hp, "PAGE")

# ---------------------------------------------------------------- utilidades de contenido
ESPACIO = {"pendiente": False}


def espaciar(p, pt=12):
    if ESPACIO["pendiente"]:
        p.paragraph_format.space_before = Pt(pt)
        ESPACIO["pendiente"] = False


MARCAS = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def runs(p, texto, tam=None):
    for trozo in MARCAS.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**"):
            r = p.add_run(trozo[2:-2]); r.bold = True
        elif trozo.startswith("`"):
            r = p.add_run(trozo[1:-1]); r.font.name = "Courier New"; r.font.size = Pt(10.5)
        elif trozo.startswith("*"):
            r = p.add_run(trozo[1:-1]); r.italic = True
        else:
            r = p.add_run(trozo)
        if tam and not trozo.startswith("`"):
            r.font.size = Pt(tam)
    return p


def P(texto, sangria=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if sangria:
        p.paragraph_format.first_line_indent = Cm(1.27)
    espaciar(p)
    return runs(p, texto)


def H1(t, nueva_pagina=True):
    h = doc.add_heading(t, level=1)
    h.paragraph_format.page_break_before = nueva_pagina
    return h


def H2(t):
    return doc.add_heading(t, level=2)


def H3(t):
    return doc.add_heading(t, level=3)


def viñetas(items, numerada=False):
    for it in items:
        p = doc.add_paragraph(style="List Number" if numerada else "List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        espaciar(p)
        runs(p, it)


def leyenda_figura(texto):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Figura "); r.italic = True
    campo(p, "SEQ Figura \\* ARABIC", "1", fmt=lambda x: setattr(x, "italic", True))
    r = p.add_run(". "); r.italic = True
    runs(p, texto)


RECORTES = RAIZ / "_herramientas" / "revision" / "recortes"


def recortar_terminal(ruta):
    """Si la imagen es una terminal (fondo oscuro al pie), quita las filas vacías del final."""
    import pymupdf
    pix = pymupdf.Pixmap(str(ruta))
    if pix.alpha:
        pix = pymupdf.Pixmap(pix, 0)
    w, h, n, m = pix.width, pix.height, pix.n, pix.samples
    def fila_oscura_vacia(y):
        base = y * pix.stride
        for x in range(4, w - 40, 3):
            o = base + x * n
            if max(m[o], m[o + 1], m[o + 2]) > 60:
                return False
        return True
    o = (h - 30) * pix.stride + (w // 2) * n
    if max(m[o], m[o + 1], m[o + 2]) > 60:          # pie claro: no es terminal (navegador, reporte)
        return ruta
    y = h - 1
    while y > 60 and fila_oscura_vacia(y):
        y -= 1
    y = min(h, y + 14)
    if y >= h - 20:
        return ruta
    RECORTES.mkdir(parents=True, exist_ok=True)
    destino = RECORTES / ruta.name
    doc_img = pymupdf.open()
    pagina = doc_img.new_page(width=w, height=y)
    pagina.insert_image(pymupdf.Rect(0, 0, w, h), pixmap=pix)
    pagina.get_pixmap(clip=pymupdf.Rect(0, 0, w, y)).save(str(destino))
    return destino


def figura(nombre, texto, ancho=16.3, carpeta=CAP):
    ruta = carpeta / nombre
    if not ruta.exists():
        raise SystemExit(f"Falta la imagen {ruta}")
    ruta = recortar_terminal(ruta)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    espaciar(p)
    p.add_run().add_picture(str(ruta), width=Cm(ancho))
    leyenda_figura(texto)


def borde(celda, lados):
    tcPr = celda._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for lado, val in lados.items():
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), "8" if val != "nil" else "0")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    tcPr.append(s)


def tabla(titulo, encabezados, filas, anchos, tam=10, col_riesgo=None, nota=None):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Tabla "); r.bold = True
    campo(p, "SEQ Tabla \\* ARABIC", "1", fmt=lambda x: setattr(x, "bold", True))
    p.add_run().add_break(WD_BREAK.LINE)
    r = p.add_run(titulo); r.italic = True
    t = doc.add_table(rows=1 + len(filas), cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for k, fila in enumerate(t.rows):
        trPr = fila._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if k == 0:
            trPr.append(OxmlElement("w:tblHeader"))
    for i, (fila, datos) in enumerate([(t.rows[0], encabezados)] + list(zip(t.rows[1:], filas))):
        for j, (celda, valor) in enumerate(zip(fila.cells, datos)):
            celda.width = Cm(anchos[j])
            celda.text = ""
            cp = celda.paragraphs[0]
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            runs(cp, str(valor), tam=tam)
            if i == 1:
                cp.paragraph_format.keep_with_next = False
            if i == 0:
                for rr in cp.runs:
                    rr.bold = True
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.keep_with_next = True
                borde(celda, {"top": "single", "bottom": "single", "left": "nil", "right": "nil"})
            else:
                ultima = i == len(filas)
                borde(celda, {"top": "nil", "bottom": "single" if ultima else "dotted", "left": "nil", "right": "nil"})
                if col_riesgo is not None and j == col_riesgo and valor in COLOR_RIESGO:
                    fondo, texto = COLOR_RIESGO[valor]
                    sombrear(celda, fondo)
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for rr in cp.runs:
                        rr.bold = True
                        rr.font.color.rgb = RGBColor.from_string(texto)
    if nota:
        n = doc.add_paragraph()
        n.paragraph_format.line_spacing = 1.0
        n.paragraph_format.space_before = Pt(4)
        n.paragraph_format.space_after = Pt(12)
        r = n.add_run("Nota. "); r.italic = True; r.font.size = Pt(10)
        runs(n, nota, tam=10)
    else:
        ESPACIO["pendiente"] = True


def codigo(texto, tam=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(10)
    espaciar(p)
    pPr = p._p.get_or_add_pPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), "F2F2F2")
    pPr.append(s)
    lineas = texto.rstrip("\n").split("\n")
    for k, l in enumerate(lineas):
        r = p.add_run(l)
        r.font.name = "Courier New"; r.font.size = Pt(tam)
        r._r.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), "Courier New")
        if k < len(lineas) - 1:
            r.add_break(WD_BREAK.LINE)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def indice(titulo, instr):
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(titulo); r.bold = True
    h.paragraph_format.page_break_before = True
    h.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    campo(p, instr, "Haga clic derecho y elija «Actualizar campos» para generar este índice.")


def leer(ruta):
    return (RAIZ / ruta).read_text(encoding="utf-8")


# ================================================================= PORTADA
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36)
p.add_run().add_picture(str(RAIZ / "_herramientas" / "recursos" / "inacap_logo.jpg"), width=Cm(6))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(36); p.paragraph_format.line_spacing = 1.15
r = p.add_run("Informe de Gestión de Vulnerabilidades"); r.bold = True; r.font.size = Pt(20)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12); p.paragraph_format.line_spacing = 1.15
r = p.add_run("Caso: Ferretería y Construcciones del Maule"); r.font.color.rgb = ROJO
r.add_break(WD_BREAK.LINE)
r = p.add_run("Ciclo de gestión de vulnerabilidades con Docker como infraestructura: "
              "aplicación web Django, base de datos Oracle, ZAP y OpenVAS en la red soc-net")
r.font.color.rgb = ROJO
datos = [("Asignatura", "Gestión de Seguridad de la Información (TI3062)"),
         ("Unidad", "Unidad 2 — Ciclo de gestión de vulnerabilidades"),
         ("Evaluación", "Sumativa 2 (ES2) — Docker como infraestructura"),
         ("Sección", "D-IEI-N6-P2-C1"),
         ("Docente", "Héctor Tejo"),
         ("Estudiantes", "Benjamín Cifuentes — Benjamín Contesso — Victor Cancino"),
         ("Fecha", "28 de octubre de 2026")]
doc.add_paragraph().paragraph_format.space_before = Pt(24)
tp = doc.add_table(rows=len(datos), cols=2)
tp.alignment = WD_TABLE_ALIGNMENT.CENTER
for fila, (k, v) in zip(tp.rows, datos):
    for celda, txt, ancho, neg in ((fila.cells[0], k, 4.0, True), (fila.cells[1], v, 8.0, False)):
        celda.width = Cm(ancho)
        celda.text = ""
        cp = celda.paragraphs[0]
        cp.paragraph_format.line_spacing = 1.15
        cp.paragraph_format.space_after = Pt(6)
        rr = cp.add_run(txt); rr.bold = neg; rr.font.size = Pt(11)

# ================================================================= ÍNDICES
indice("Índice de contenido", 'TOC \\o "1-3" \\h \\z \\u')
indice("Índice de figuras", 'TOC \\h \\z \\c "Figura"')
indice("Índice de tablas", 'TOC \\h \\z \\c "Tabla"')

# ================================================================= 1. INTRODUCCIÓN
H1("1. Introducción")
P("El presente informe documentó la aplicación del **ciclo de gestión de vulnerabilidades** sobre una "
  "infraestructura construida completamente con contenedores Docker. El ciclo, que el docente presentó como "
  "basado en el marco del NIST, propone identificar las vulnerabilidades de una infraestructura, evaluarlas y "
  "priorizarlas según su impacto, actuar sobre ellas (corregir, mitigar o aceptar el riesgo), verificar que las "
  "correcciones funcionaron y mejorar el proceso de forma continua (Souppaya y Scarfone, 2022; CrowdStrike, s.f.).")
P("Como caso se retomó la empresa *Ferretería y Construcciones del Maule*, analizada en la Evaluación 1. Se "
  "desarrolló una aplicación web propia en Django que publica el catálogo de productos de la ferretería y lo lee "
  "desde una base de datos Oracle. Ambos componentes, junto con dos escáneres de seguridad —ZAP para "
  "aplicaciones web y OpenVAS para infraestructura—, se desplegaron como cuatro contenedores conectados a una red "
  "Docker externa llamada `soc-net`. Cada contenedor quedó definido en su propio archivo `docker-compose.yml`, de "
  "modo que su configuración queda guardada y se puede recrear sin depender de comandos sueltos.")
P("El documento presenta primero el marco teórico (el ciclo y los conceptos de Docker solicitados) y el trabajo "
  "previo (alcance, herramientas e inventario de activos). Luego relata paso a paso la construcción de la "
  "infraestructura, incluidos los dos problemas reales que aparecieron y cómo se resolvieron, los cuatro escaneos "
  "exigidos (dos con ZAP y dos pasadas de OpenVAS), la tabla de resumen de vulnerabilidades y la estrategia de "
  "mitigación priorizada. Finalmente se aplicaron tres correcciones y se volvió a escanear la aplicación, cerrando "
  "una vuelta completa del ciclo. Todo el proceso se respaldó con capturas tomadas durante la ejecución real.")

# ================================================================= 2. OBJETIVOS
H1("2. Objetivos")
H2("2.1 Objetivo general")
P("Aplicar el ciclo de gestión de vulnerabilidades a una infraestructura de contenedores Docker compuesta por una "
  "aplicación web propia, una base de datos Oracle y dos escáneres de seguridad, documentando el proceso con "
  "evidencias que permitan repetirlo.")
H2("2.2 Objetivos específicos")
viñetas([
    "Construir la imagen de la aplicación web a partir de un Dockerfile y desplegar los cuatro contenedores en la red externa `soc-net`, cada uno con su archivo `docker-compose.yml`.",
    "Identificar vulnerabilidades en la aplicación web mediante dos escaneos con ZAP y en su contenedor mediante dos pasadas con OpenVAS.",
    "Evaluar y priorizar los hallazgos en una tabla de resumen que indique el activo afectado, el impacto y el riesgo según la categoría de cada herramienta.",
    "Proponer una estrategia de mitigación ordenada por prioridad y verificar con un nuevo escaneo el efecto de las correcciones aplicadas.",
], numerada=False)

# ================================================================= 3. MARCO TEÓRICO
H1("3. Marco teórico")
H2("3.1 El ciclo de gestión de vulnerabilidades")
P("La gestión de vulnerabilidades no es una revisión única, sino un proceso **cíclico**: una aplicación que hoy no "
  "presenta vulnerabilidades puede tenerlas mañana, por ejemplo, cuando se publica una falla en la versión de "
  "Python o de una librería que utiliza. La guía SP 800-40 del NIST plantea la gestión de parches y "
  "vulnerabilidades como mantenimiento preventivo permanente (Souppaya y Scarfone, 2022). La Figura 1 muestra el "
  "ciclo entregado en el enunciado, que coincide con el ciclo de vida descrito por CrowdStrike (s.f.): un trabajo "
  "previo y cinco fases que se repiten.")
figura("enunciado-ciclo.png", "Ciclo de gestión de vulnerabilidades. Tomado del enunciado de la Evaluación 2 (U2ES2).", 14, DIA)
P("La Tabla 1 resume cómo se aplicó cada fase en este trabajo.")
tabla("Fases del ciclo de gestión de vulnerabilidades aplicadas en la evaluación",
      ["Fase", "Actividades (según la figura)", "Aplicación en este trabajo", "Sección"],
      [["Trabajo previo", "Determinar el alcance; definir funciones y responsabilidades; seleccionar herramientas; preparar el inventario de activos", "Alcance: 4 contenedores en `soc-net`. Herramientas: ZAP y OpenVAS. Inventario de activos con IP y puertos", "4"],
       ["Identificar", "Identificar; escanear", "Construcción de la infraestructura; 2 escaneos con ZAP, 2 pasadas con OpenVAS y revisión de configuración", "5 y 6"],
       ["Evaluar y priorizar", "Agregar contexto de amenaza; asignar valor", "Tabla de resumen con activo, impacto y riesgo de cada hallazgo", "7"],
       ["Actuar", "Corregir; mitigar; aceptar el riesgo", "Estrategia de mitigación priorizada", "8"],
       ["Verificar", "Validar; volver a escanear", "Aplicación de 3 correcciones (app 1.1) y escaneo de verificación con ZAP", "9"],
       ["Mejorar", "Evaluar métricas; automatizar procesos", "Métricas de los escaneos y propuesta de escaneos periódicos", "9"]],
      [2.6, 4.6, 6.6, 1.6])

H2("3.2 Conceptos de Docker")
P("Los siguientes conceptos, solicitados en el enunciado, se definen a partir de la documentación oficial de Docker.")
H3("Docker y contenedor")
P("Docker es una plataforma abierta para desarrollar, distribuir y ejecutar aplicaciones dentro de "
  "**contenedores**. Un contenedor es una instancia en ejecución de una imagen: un proceso aislado que incluye todo "
  "lo que la aplicación necesita (código, intérprete, librerías y configuración) y que comparte el núcleo del "
  "sistema anfitrión. Esto permite separar la aplicación de la infraestructura y ejecutar en un mismo equipo, por "
  "ejemplo, dos aplicaciones con versiones distintas de Python sin que sus librerías se mezclen (Docker Inc., s.f.-a).")
H3("Imagen (image)")
P("Una imagen es una plantilla de solo lectura con las instrucciones para crear un contenedor. Está formada por "
  "capas, se identifica con un nombre y una etiqueta (`nombre:tag`, donde la etiqueta suele indicar la versión) y "
  "se obtiene desde un registro como Docker Hub o se construye a partir de un Dockerfile (Docker Inc., s.f.-a). En "
  "este trabajo, Oracle, ZAP y OpenVAS usan imágenes publicadas; solo la imagen de la aplicación web "
  "(`ferreteria-web`) es propia.")
H3("Dockerfile")
P("Un Dockerfile es un archivo de texto que contiene, en orden, las instrucciones (`FROM`, `WORKDIR`, `COPY`, "
  "`RUN`, `EXPOSE`, `CMD`, entre otras) que el comando `docker build` ejecuta para construir una imagen (Docker "
  "Inc., s.f.-d). El detalle de cada instrucción se desarrolla en el documento complementario `Dockerfile.md`.")
H3("Volumen (volume)")
P("Un volumen es un mecanismo de almacenamiento persistente administrado por Docker e independiente del ciclo de "
  "vida del contenedor: los datos se conservan aunque el contenedor se elimine o se vuelva a crear (Docker Inc., "
  "s.f.-c). Se usaron los volúmenes `oracle-volume` (datos de Oracle) y `openvas-data` (base de datos, feeds y "
  "resultados de OpenVAS).")
H3("Redes en Docker")
P("Las redes permiten que los contenedores se comuniquen entre sí y con el exterior. El controlador por defecto "
  "es `bridge`; una red *bridge* definida por el usuario, como `soc-net`, agrupa a los contenedores, los aísla de "
  "otros y permite que se encuentren por nombre gracias al DNS interno de Docker (Docker Inc., s.f.-b). Dos "
  "contenedores en redes distintas no se comunican, lo que constituye una primera estrategia de seguridad. En "
  "Compose, `external: true` indica que la red ya existe y fue creada antes, fuera del archivo.")
H3("Docker Compose")
P("Docker Compose es la herramienta para definir y ejecutar aplicaciones de uno o varios contenedores mediante un "
  "archivo YAML (`docker-compose.yml`) que declara servicios, imagen, puertos, redes, volúmenes y variables. Se "
  "levanta con `docker compose up -d` (Docker Inc., s.f.-e). Su ventaja frente a un comando suelto es que la "
  "configuración queda guardada en un archivo versionable, por lo que el contenedor se puede recrear siempre igual.")

H2("3.3 Herramientas de evaluación de vulnerabilidades")
P("**ZAP** (Zed Attack Proxy) es un escáner de seguridad para aplicaciones web de código abierto. Explora el sitio "
  "con su *spider*, analiza pasivamente cada respuesta y, en el escaneo activo, envía peticiones de ataque "
  "(inyección SQL, XSS, entre otras) para confirmar vulnerabilidades. Clasifica sus alertas en Alto, Medio, Bajo e "
  "Informativo. Se usó su imagen oficial con la interfaz gráfica completa servida en el navegador mediante "
  "Webswing (ZAP, s.f.-a; ZAP, s.f.-b).")
P("**OpenVAS**, parte de Greenbone Community Edition, es un escáner de vulnerabilidades de infraestructura de "
  "código abierto, alternativa libre a Nessus. Ejecuta decenas de miles de pruebas (NVT) contra los servicios de un "
  "equipo y asigna a cada hallazgo una severidad CVSS (Greenbone AG, s.f.). Se usó la imagen todo-en-uno de "
  "Immauss, que incluye el escáner, el gestor y la interfaz web GSA (Immauss, s.f.).")

# ================================================================= 4. TRABAJO PREVIO
H1("4. Trabajo previo")
H2("4.1 Alcance")
P("El alcance de la evaluación son los cuatro contenedores definidos en el enunciado y conectados a la red "
  "`soc-net` (Figura 2). Los objetivos de los escaneos son la aplicación web (ZAP) y el contenedor que la aloja "
  "(OpenVAS). Además, se revisó la configuración de los cuatro archivos `docker-compose.yml` en busca de malas "
  "prácticas.")
figura("enunciado-arquitectura.png", "Arquitectura solicitada: cuatro contenedores en la red soc-net. Tomado del enunciado (U2ES2).", 12, DIA)
H2("4.2 Funciones y responsabilidades")
tabla("Funciones definidas para el proceso",
      ["Función", "Responsabilidad"],
      [["Responsable de infraestructura", "Crear la red, escribir los archivos Compose y el Dockerfile, construir la imagen y mantener los contenedores en ejecución"],
       ["Analista de vulnerabilidades", "Configurar ZAP y OpenVAS, ejecutar los escaneos y exportar los reportes"],
       ["Responsable de remediación", "Analizar los hallazgos, priorizarlos, aplicar las correcciones y solicitar el escaneo de verificación"],
       ["Documentación", "Registrar cada paso con capturas, redactar el informe y preparar la demostración"]],
      [5.0, 10.4])
H2("4.3 Herramientas seleccionadas")
tabla("Herramientas y versiones utilizadas",
      ["Herramienta", "Versión", "Uso"],
      [["Docker Desktop / Docker Engine", "Engine 29.7.2, Compose v5.3.1", "Plataforma de contenedores (Windows 11 + WSL 2)"],
       ["Oracle Database Free", "gvenzl/oracle-free:latest (26ai)", "Base de datos de la aplicación"],
       ["Python / Django", "python:3.12-slim / Django 5.2.18 LTS", "Aplicación web (catálogo de la ferretería)"],
       ["ZAP", "zaproxy/zap-stable 2.17.0", "Escáner de aplicaciones web"],
       ["OpenVAS (Greenbone)", "immauss/openvas 26.07.27.01 (GSA 27.5.0, gvmd 26.34.0)", "Escáner de infraestructura"],
       ["YAMLTools Visualizer", "yamltools.dev (web)", "Diagramas de los archivos Compose"],
       ["Asistente de IA (Claude)", "—", "Apoyo para generar y explicar archivos, resumir reportes y redactar; el grupo revisó cada resultado"]],
      [4.2, 5.4, 5.8])
H2("4.4 Inventario de activos")
P("La Tabla 4 presenta el inventario de activos. Las direcciones IP se obtuvieron con el comando de red que se "
  "muestra en la sección 5.6.")
tabla("Inventario de activos de la infraestructura",
      ["Activo", "Imagen", "IP en soc-net", "Puertos (host:contenedor)", "Función", "Criticidad"],
      [["oracle-db", "gvenzl/oracle-free:latest", "172.28.0.3", "1521:1521", "Base de datos (tablas CATEGORIA y PRODUCTO)", "Alta"],
       ["app-web", "ferreteria-web:1.0 → 1.1", "172.28.0.20 (fija)", "8088:8000", "Catálogo web y panel de administración", "Alta"],
       ["zap", "zaproxy/zap-stable:latest", "172.28.0.4", "8080:8080, 8090:8090", "Escáner web (interfaz y API)", "Media"],
       ["openvas", "immauss/openvas:latest", "172.28.0.2", "9392:9392", "Escáner de infraestructura", "Media"],
       ["soc-net", "Red bridge externa", "172.28.0.0/16 (gw .1)", "—", "Comunicación entre contenedores", "Alta"],
       ["Volúmenes", "oracle-volume, openvas-data", "—", "—", "Persistencia de datos", "Alta"]],
      [2.2, 3.6, 2.6, 2.6, 3.0, 1.6], tam=9)

# ================================================================= 5. DESARROLLO
H1("5. Construcción de la infraestructura")
P("Esta sección relata, en el orden en que se ejecutó el 27 de octubre de 2026, cómo se construyó la "
  "infraestructura. OpenVAS se levantó en segundo lugar, inmediatamente después de la red, porque al iniciar debe "
  "cargar su base de pruebas de vulnerabilidades y el docente advirtió que puede tardar entre 10 minutos y una "
  "hora. Todos los archivos se encuentran en las carpetas `01-oracle`, `02-app-django`, `03-zap` y `04-openvas`.")
H2("5.1 Red externa soc-net")
P("Primero se revisaron las redes existentes (Figura 3). La red `soc-net` existía de una prueba anterior con otra "
  "subred, por lo que se eliminó y se volvió a crear con la subred `172.28.0.0/16` indicada en el archivo de "
  "configuración del curso (Figuras 4 y 5). Al ser externa, debe existir antes de levantar cualquier contenedor.")
figura("01-redes-estado-inicial.png", "Redes Docker existentes antes de comenzar (`docker network ls`).", 14)
figura("02-crear-red-soc-net.png", "Creación de la red externa soc-net con la subred 172.28.0.0/16.")
figura("03-inspeccion-red-soc-net.png", "Inspección de soc-net: controlador bridge, subred 172.28.0.0/16 y gateway 172.28.0.1.")

H2("5.2 Contenedor OpenVAS")
P("El archivo `04-openvas/docker-compose.yml` (Figura 6) usa la imagen `immauss/openvas`, publica la interfaz web "
  "en el puerto 9392, guarda los datos en el volumen `openvas-data` y conecta el contenedor a `soc-net`. La "
  "contraseña del usuario `admin` no se escribe en el archivo: se lee de la variable `OPENVAS_PASSWORD` definida en "
  "un archivo `.env` que no se entrega. El contenedor se levantó con `docker compose up -d` (Figura 7) y en sus "
  "registros se observa la restauración de la base de datos inicial (Figura 8).")
figura("04-compose-openvas.png", "Archivo docker-compose.yml de OpenVAS (versión inicial).")
figura("05-openvas-compose-up.png", "Levantamiento de OpenVAS: se crea el volumen y se inicia el contenedor.")
figura("06-openvas-logs-arranque.png", "Registros de arranque de OpenVAS: certificados y restauración de la base de datos.")
H3("Incidente: reinicios en bucle por la sincronización de feeds")
P("A los 25 minutos el gestor de OpenVAS aún no respondía y el contenedor llevaba dos reinicios. Al revisar los "
  "registros (Figura 9) se encontró la causa: en cada arranque, el script de la imagen intenta descargar los feeds "
  "desde el servidor público de Greenbone mediante `rsync`; la conexión se cortaba (`rsync error ... code 12`) y, "
  "según el propio script, si la sincronización falla el contenedor termina, Docker lo reinicia y el ciclo se "
  "repite. La solución fue agregar la variable `SKIPSYNC: \"true\"`, documentada por la imagen, para usar los feeds "
  "incluidos en ella (fechados el 6 de octubre de 2026), y recrear el contenedor conservando el volumen "
  "(Figura 10). Tres minutos después OpenVAS quedó operativo, con la configuración de escaneo «Full and fast» y "
  "más de 192 000 pruebas (NVT) cargadas. Esta decisión deja un riesgo que se registra como hallazgo V14: los "
  "feeds no incluyen vulnerabilidades publicadas después de esa fecha.")
figura("06b-openvas-error-sincronizacion.png", "Causa de los reinicios: la descarga de feeds por rsync se corta y el contenedor se detiene.")
figura("06c-openvas-skipsync-recrear.png", "Corrección: variable SKIPSYNC en el compose y recreación del contenedor.")

H2("5.3 Contenedor Oracle")
P("Para Oracle se utilizó, sin modificaciones, el archivo `docker-compose.yml` entregado por el docente en el "
  "material del curso (Figura 11). Este archivo contiene las contraseñas escritas directamente en las variables de "
  "entorno; el docente explicó en clase que lo dejó así a propósito, como mala práctica, para ver si los escáneres "
  "la detectan. Al levantarlo, Docker advierte que el atributo `version` está obsoleto (Figura 12), lo que "
  "coincide con lo señalado en clase: esa línea es opcional. La base de datos quedó lista en 135 segundos "
  "(Figura 13).")
figura("07-compose-oracle.png", "Archivo docker-compose.yml de Oracle entregado por el docente (sin cambios).")
figura("08-oracle-compose-up.png", "Levantamiento de Oracle con docker compose up -d.")
figura("09-oracle-lista.png", "Registros de Oracle: «DATABASE IS READY TO USE!».")
H3("Incidente: el usuario de aplicación no se creó")
P("Al intentar crear las tablas con el usuario `usuario`, Oracle rechazó la conexión con el error `ORA-01017` "
  "(Figura 14). Los registros mostraron la causa (Figura 15): el compose pide crear la base de datos acoplable "
  "`FREEPDB1` (`ORACLE_DATABASE: FREEPDB1`), pero la versión actual de la imagen ya la incluye; el script de "
  "inicio falló con `ORA-65012`, el contenedor se reinició y en el segundo arranque se consideró ya inicializado, "
  "por lo que nunca creó el usuario de aplicación. Sin modificar el archivo del docente, se creó el usuario con el "
  "script `createAppUser` que trae la imagen (Figura 16).")
figura("10a-oracle-error-ora01017.png", "Primer intento de crear las tablas: error ORA-01017 (credenciales inválidas).")
figura("10b-oracle-causa-ora65012.png", "Causa: ORA-65012 en el primer arranque, reinicio del contenedor y omisión del usuario de aplicación.")
figura("10c-oracle-crear-usuario-app.png", "Creación del usuario de aplicación con createAppUser y verificación en DBA_USERS.")
P("Con el usuario disponible se ejecutó el script `01-oracle/sql/01_tablas.sql`, que crea las tablas "
  "`CATEGORIA(id, nombre)` y `PRODUCTO(codigo, nombre, precio, stock, unidad, categoriaid)` con dos registros cada "
  "una (Figura 17), y se verificó su contenido, incluida una consulta que une ambas tablas (Figura 18).")
figura("10-oracle-crear-tablas.png", "Creación de las tablas CATEGORIA y PRODUCTO e inserción de dos registros en cada una.")
figura("11-oracle-consultar-tablas.png", "Consulta de las tablas y unión de producto con su categoría.")

H2("5.4 Aplicación web con Dockerfile")
P("La aplicación es un catálogo de la ferretería desarrollado en Python 3.12 con Django 5.2 (versión de soporte "
  "extendido). Lista los productos con su categoría, permite buscarlos por nombre, muestra el detalle de cada uno, "
  "incluye el panel de administración de Django y una ruta `/salud/` que confirma la conexión con Oracle. Se "
  "conecta a la base de datos con el driver `oracledb` en modo *Thin*, que no requiere instalar Oracle Instant "
  "Client (Oracle, s.f.), usando la cadena `oracle-db:1521/FREEPDB1`: `oracle-db` es el nombre del contenedor, "
  "que el DNS de `soc-net` resuelve. La Figura 19 muestra la estructura del proyecto.")
figura("12-estructura-app-django.png", "Estructura del proyecto Django (carpeta 02-app-django).", 11)
P("El Dockerfile (Figura 20) parte de la imagen oficial `python:3.12-slim`, instala las dependencias en una capa "
  "propia para aprovechar la caché, copia el código (el archivo `.dockerignore` excluye el `.env` con los "
  "secretos), reúne los archivos estáticos, crea un usuario sin privilegios para no ejecutar la aplicación como "
  "*root*, documenta el puerto 8000 y define un chequeo de salud. Su `ENTRYPOINT` ejecuta un script que espera a "
  "que Oracle acepte conexiones —necesario porque Oracle se levanta con otro archivo Compose y `depends_on` no "
  "aplica— y aplica las migraciones de Django; luego cede el control al `CMD`, que inicia el servidor gunicorn. "
  "Cada instrucción se explica en detalle en `Dockerfile.md`.")
figura("13-dockerfile-app.png", "Dockerfile comentado de la aplicación web.")
P("La imagen se construyó con `docker build -t ferreteria-web:1.0 .` en 79 segundos (Figura 21) y quedó "
  "disponible junto a las imágenes descargadas (Figura 22). Esta es la única imagen que no existe en Docker Hub.")
figura("14-docker-build-app.png", "Construcción de la imagen ferreteria-web:1.0 (12 de 12 pasos).")
figura("15-docker-images.png", "Imágenes disponibles: la imagen propia y las descargadas.")
P("El archivo `02-app-django/docker-compose.yml` (Figura 23) publica el puerto 8000 del contenedor en el 8088 del "
  "equipo, carga las variables desde `.env` y asigna la IP fija `172.28.0.20` en `soc-net`, para que los "
  "escáneres tengan un objetivo estable. Al levantarla (Figura 24), los registros muestran la conexión con Oracle, "
  "las migraciones de Django, la creación del superusuario y el inicio de gunicorn (Figura 25).")
figura("16-compose-app.png", "Archivo docker-compose.yml de la aplicación web.")
figura("17-app-compose-up.png", "Levantamiento del contenedor app-web.")
figura("18-app-logs-arranque.png", "Registros de app-web: Oracle disponible, migraciones, superusuario y gunicorn.")
P("Las Figuras 26 a 30 muestran la aplicación funcionando en `http://localhost:8088`: el listado leído desde "
  "Oracle, la búsqueda, el ingreso al panel de administración y la gestión de productos.")
figura("19-app-listado-productos.png", "Catálogo de productos leído desde Oracle (http://localhost:8088).", 14)
figura("20-app-busqueda.png", "Búsqueda de productos por nombre.", 14)
figura("21a-app-admin-login.png", "Formulario de ingreso al panel de administración (/admin/).", 13)
figura("21b-app-admin-panel.png", "Panel de administración con los modelos del catálogo.", 14)
figura("21c-app-admin-productos.png", "Productos de Oracle administrados desde el panel de Django.", 14)

H2("5.5 Contenedor ZAP")
P("El archivo `03-zap/docker-compose.yml` (Figura 31) usa la imagen oficial `zaproxy/zap-stable` (ZAP 2.17.0, la "
  "misma versión usada por el docente) con el comando `zap-webswing.sh`, que ofrece la interfaz gráfica completa "
  "en el navegador. Publica la interfaz en el puerto 8080 y el proxy/API en el 8090; la API queda protegida con "
  "una clave leída desde `.env`, y los reportes se guardan en la carpeta `reportes/zap` del proyecto mediante un "
  "volumen. La Figura 33 muestra la interfaz en `http://localhost:8080/zap/`.")
figura("22-compose-zap.png", "Archivo docker-compose.yml de ZAP.")
figura("23-zap-compose-up.png", "Levantamiento del contenedor ZAP.")
figura("24-zap-interfaz-web.png", "Interfaz de ZAP 2.17.0 en el navegador (Webswing).")

H2("5.6 Evidencia de la red: contenedores e IP")
P("Con los cuatro contenedores en ejecución (Figura 34) se ejecutó el comando entregado por el docente para ver "
  "los contenedores de `soc-net` con su IP (Figura 35). En PowerShell de Windows las comillas internas `\"\\n\"` "
  "deben escaparse como `\\\"\\n\\\"` (o reemplazarse por `{{println}}`); el resultado es el mismo. Estas IP son "
  "las que se entregan a los escáneres. La Figura 36 muestra los volúmenes creados.")
figura("25-docker-ps-4-contenedores.png", "Los cuatro contenedores en ejecución (docker ps).")
figura("26-red-soc-net-ip-contenedores.png", "Contenedores de soc-net con su dirección IP (comando del docente).")
figura("27-docker-volumenes.png", "Volúmenes persistentes de Oracle y OpenVAS.", 12)

H2("5.7 Diagramas de los archivos Compose")
P("Cada archivo Compose se cargó en el visualizador de YAMLTools, reemplazando antes las contraseñas por "
  "`\"[oculto]\"` para no enviarlas a un sitio externo. Las Figuras 37 a 40 muestran su estructura en la vista "
  "*Diagrama de Bloques*: servicios, volúmenes y la red externa `soc-net`.")
figura("diagrama-01-oracle.png", "Diagrama de bloques del compose de Oracle (YAMLTools).", 13, DIA)
figura("diagrama-02-app-django.png", "Diagrama de bloques del compose de la aplicación web (YAMLTools).", 13, DIA)
figura("diagrama-03-zap.png", "Diagrama de bloques del compose de ZAP (YAMLTools).", 13, DIA)
figura("diagrama-04-openvas.png", "Diagrama de bloques del compose de OpenVAS (YAMLTools).", 13, DIA)

# ================================================================= 6. ESCANEOS
H1("6. Identificación: escaneos de vulnerabilidades")
P("Los escaneos se lanzaron desde los contenedores de seguridad hacia `http://172.28.0.20:8000`, es decir, la IP "
  "de la aplicación dentro de `soc-net` y su **puerto interno** (8000), no el 8088 publicado en el equipo. ZAP se "
  "operó mediante su API y su interfaz web muestra el avance en vivo; OpenVAS se configuró por su protocolo de "
  "gestión (GMP) y se revisó en su interfaz GSA.")
H2("6.1 ZAP — Escaneo 1: exploración y análisis pasivo")
P("El primer escaneo ejecutó el *spider* de ZAP, que recorrió el sitio y encontró 16 URL en 11 segundos "
  "(Figura 41), y el análisis pasivo de todas las respuestas. Se obtuvieron dos alertas de riesgo Medio, una Baja "
  "y tres Informativas (Figura 42). El resultado se exportó como `Reporte1.json` y `Reporte1.html` (Figura 43).")
figura("28-zap-escaneo1-spider.png", "Escaneo 1: árbol del sitio descubierto por el spider de ZAP.")
figura("29-zap-escaneo1-alertas.png", "Escaneo 1: alertas de ZAP y detalle de «Content Security Policy (CSP) Header Not Set».")
figura("30-zap-reporte1-html.png", "Reporte 1 de ZAP en HTML: resumen de alertas por nivel de riesgo.", 14)
H2("6.2 ZAP — Escaneo 2: escaneo activo")
P("El segundo escaneo fue activo: ZAP envió 3 444 peticiones con pruebas de ataque durante cerca de cinco minutos "
  "y medio (Figuras 44 y 45). Se intentaron, entre otras, inyecciones SQL y XSS sobre el parámetro de búsqueda "
  "`q` y el formulario de ingreso; ninguna se confirmó, y los intentos de envío al formulario sin token fueron "
  "rechazados por la protección CSRF de Django (código 403). Se agregaron las alertas «HTTP Only Site» (Medio) y "
  "«User Agent Fuzzer» (Informativa) (Figura 46). El resultado se exportó como `Reporte2.json` y `Reporte2.html` "
  "(Figuras 47 y 48).")
figura("31-zap-escaneo2-activo-progreso.png", "Escaneo 2: escaneo activo en curso (35 %).")
figura("36-zap-escaneo2-activo-completo.png", "Escaneo 2: escaneo activo terminado (3 444 peticiones).")
figura("37-zap-escaneo2-alertas.png", "Escaneo 2: alerta «HTTP Only Site» detectada por el escaneo activo.")
figura("38-zap-reporte2-html.png", "Reporte 2 de ZAP en HTML.", 14)
figura("39-zap-reportes-generados.png", "Reportes de ZAP generados en formato JSON y HTML.", 13)
tabla("Resumen de alertas de ZAP por escaneo",
      ["Alerta (regla ZAP)", "CWE", "Escaneo 1", "Escaneo 2", "Riesgo"],
      [["Content Security Policy (CSP) Header Not Set (10038)", "693", "Sí", "Sí", "Medio"],
       ["Cross-Domain Misconfiguration (10098)", "264", "Sí", "Sí", "Medio"],
       ["HTTP Only Site (10106)", "311", "—", "Sí", "Medio"],
       ["Cookie No HttpOnly Flag (10010)", "1004", "Sí", "Sí", "Bajo"],
       ["Authentication Request Identified (10111)", "—", "Sí", "Sí", "Informativo"],
       ["Session Management Response Identified (10112)", "—", "Sí", "Sí", "Informativo"],
       ["User Controllable HTML Element Attribute (10031)", "20", "Sí", "Sí", "Informativo"],
       ["User Agent Fuzzer (10104)", "—", "—", "Sí", "Informativo"]],
      [7.2, 1.4, 2.0, 2.0, 2.6], col_riesgo=4,
      nota="Riesgo según la clasificación de ZAP. Fuente: Reporte1.json y Reporte2.json.")

H2("6.3 OpenVAS — Pasada 1")
P("Se ingresó a la interfaz GSA en `http://localhost:9392` (Figura 49); el panel muestra las más de 194 000 NVT "
  "disponibles (Figura 50). Se creó el objetivo `app-web Ferreteria (172.28.0.20)` con la lista de puertos *All "
  "IANA assigned TCP* y la prueba de disponibilidad *Consider Alive* (Figura 51), y una tarea con la configuración "
  "*Full and fast* (Figura 52). La primera pasada duró cinco minutos y encontró dos hallazgos de severidad Baja y "
  "once registros informativos (Figuras 53 y 54). Se exportó en XML, CSV y PDF.")
figura("32-openvas-login.png", "Ingreso a la interfaz web de OpenVAS (GSA). Se observa la advertencia de conexión sin cifrar.", 14)
figura("33-openvas-dashboard.png", "Panel de OpenVAS con las NVT cargadas por severidad.", 14)
figura("34-openvas-objetivo.png", "Objetivo creado: contenedor de la aplicación web (172.28.0.20).", 14)
figura("35-openvas-tarea-creada.png", "Tarea «Full and fast» creada sobre el objetivo.", 14)
figura("40a-openvas-pasada1-informacion.png", "Pasada 1 terminada: duración, estado y cantidad de resultados.", 14)
figura("40b-openvas-pasada1-resultados.png", "Pasada 1: hallazgos de severidad Baja.", 14)
H2("6.4 OpenVAS — Pasada 2")
P("El docente pidió dos pasadas porque una sola no siempre es confiable. La segunda pasada, sobre la misma tarea, "
  "duró seis minutos y obtuvo exactamente los mismos resultados (Figuras 55 a 57), lo que confirma que los "
  "hallazgos son consistentes; la tendencia de la tarea aparece estable («→»). La Figura 58 muestra los reportes "
  "exportados y la Tabla 6 compara ambas pasadas.")
figura("41-openvas-tarea-dos-pasadas.png", "Tarea con dos reportes (dos pasadas), estado Done y severidad 2,6 (Baja).", 14)
figura("42-openvas-reportes-dos-pasadas.png", "Reportes de ambas pasadas con el conteo por severidad.", 14)
figura("43-openvas-pasada2-resultados.png", "Pasada 2: mismos hallazgos de severidad Baja.", 14)
figura("44-openvas-reportes-exportados.png", "Reportes de OpenVAS exportados (XML, CSV y PDF) por pasada.", 13)
tabla("Comparación de las dos pasadas de OpenVAS",
      ["Hallazgo (NVT)", "CVSS", "Pasada 1", "Pasada 2", "Riesgo"],
      [["TCP Timestamps Information Disclosure", "2,6", "Sí", "Sí", "Bajo"],
       ["ICMP Timestamp Reply Information Disclosure", "2,1", "Sí", "Sí", "Bajo"],
       ["Gunicorn Detection (HTTP) / HTTP Server Banner", "0,0", "Sí", "Sí", "Informativo"],
       ["Otros registros (detección de SO, servicios, cabeceras, CPE, traceroute)", "0,0", "Sí (9)", "Sí (9)", "Informativo"]],
      [7.4, 1.4, 2.0, 2.0, 2.4], col_riesgo=4,
      nota="Pasada 1: 07:23–07:29; pasada 2: 07:31–07:36 (hora de Chile). Fuente: OpenVAS-pasada1.xml y OpenVAS-pasada2.xml.")
H2("6.5 Revisión de configuración")
P("Además de los escáneres, se revisaron los archivos Compose, los registros y las versiones utilizadas. Esta "
  "revisión encontró la mala práctica que el docente dejó a propósito —las credenciales de Oracle en texto plano—, "
  "que ninguno de los dos escáneres puede detectar porque no analizan archivos de configuración, y otros hallazgos "
  "que se incorporan a la tabla de resumen: el puerto 1521 publicado en todas las interfaces del equipo, la "
  "contraseña débil y repetida, las imágenes sin versión fija (`latest`), la clave de la API de ZAP visible en "
  "`docker logs zap`, la interfaz de OpenVAS sin cifrar y el driver `oracledb` 26.0.1, que está fuera del rango "
  "soportado oficialmente por Django 5.2 (2.3.0 a 4.0.2) (Django Software Foundation, s.f.-a).")

# ================================================================= 7. TABLA DE VULNERABILIDADES
H1("7. Evaluación y priorización: tabla de resumen de vulnerabilidades")
P("La Tabla 7 consolida los hallazgos únicos. Para ZAP y OpenVAS el riesgo corresponde a la categoría que asigna "
  "cada herramienta, como indicó el docente; para los hallazgos de la revisión de configuración el grupo asignó la "
  "categoría considerando el impacto sobre la aplicación. Los colores siguen el formato del enunciado.")
V = [
    ["V01", "Credenciales de Oracle en texto plano en el docker-compose.yml (ORACLE_PASSWORD, APP_USER_PASSWORD)", "Revisión de configuración", "oracle-db (compose)", "Quien acceda al archivo obtiene acceso total a la base de datos (SYS y SYSTEM)", "Alto"],
    ["V02", "Contraseña débil y repetida (MyPass123) para SYS, SYSTEM y el usuario de aplicación", "Revisión de configuración", "oracle-db", "Fácil de adivinar; si se compromete una cuenta, se comprometen todas", "Alto"],
    ["V03", "Puerto 1521 de Oracle publicado en todas las interfaces del equipo (0.0.0.0:1521)", "Revisión (docker ps)", "oracle-db", "La base de datos queda expuesta a la red local además de soc-net", "Medio"],
    ["V04", "Sitio servido solo por HTTP, sin HTTPS", "ZAP esc. 2 (10106)", "app-web", "Credenciales del panel y cookies viajan sin cifrar y pueden interceptarse", "Medio"],
    ["V05", "Cabecera Content-Security-Policy no configurada", "ZAP esc. 1 y 2 (10038)", "app-web", "Sin una segunda barrera ante XSS e inyección de contenido", "Medio"],
    ["V06", "CORS permisivo en archivos estáticos (Access-Control-Allow-Origin: *)", "ZAP esc. 1 y 2 (10098)", "app-web (/static/)", "Cualquier sitio puede leer los estáticos; impacto real bajo porque son públicos", "Medio"],
    ["V07", "Interfaz web de OpenVAS sin cifrar (HTTP)", "Revisión (login GSA)", "openvas", "La contraseña del administrador del escáner viaja en claro", "Medio"],
    ["V08", "Cookie csrftoken sin atributo HttpOnly", "ZAP (10010)", "app-web", "JavaScript puede leer el token CSRF; Django indica que no aporta protección práctica", "Bajo"],
    ["V09", "TCP Timestamps Information Disclosure", "OpenVAS pasadas 1 y 2", "app-web (172.28.0.20)", "Permite estimar el tiempo de actividad del sistema", "Bajo"],
    ["V10", "ICMP Timestamp Reply Information Disclosure", "OpenVAS pasadas 1 y 2", "app-web (172.28.0.20)", "Revela la hora del sistema y facilita el reconocimiento", "Bajo"],
    ["V11", "Driver oracledb 26.0.1 fuera del rango soportado por Django 5.2", "Revisión de dependencias", "imagen ferreteria-web:1.0", "Combinación no probada por Django; riesgo de fallas o parches no coordinados", "Bajo"],
    ["V12", "Imágenes con etiqueta latest (oracle-free, zap-stable, openvas)", "Revisión de configuración", "composes 01, 03 y 04", "Versiones no reproducibles; no se sabe con certeza qué versión se evaluó", "Bajo"],
    ["V13", "Clave de la API de ZAP visible en docker logs zap", "Revisión (docker logs)", "zap", "Quien lea los registros puede controlar ZAP por su API", "Bajo"],
    ["V14", "Feeds de OpenVAS sin actualizar (SKIPSYNC, feeds del 06-10-2026)", "Incidente de despliegue", "openvas", "No se detectan vulnerabilidades publicadas después de esa fecha", "Bajo"],
    ["V15", "Divulgación de la tecnología del servidor (Server: gunicorn)", "OpenVAS (registro) / cabeceras", "app-web", "Facilita el reconocimiento; no expone la versión", "Informativo"],
    ["V16", "Atributo HTML controlable por el usuario (posible XSS)", "ZAP (10031)", "app-web (/?q, login)", "Django escapa la salida; el escaneo activo no confirmó XSS", "Informativo"],
    ["V17", "Solicitud de autenticación y gestión de sesión identificadas", "ZAP (10111, 10112)", "app-web (/admin/login/)", "Solo identifica el formulario de ingreso y las cookies de sesión", "Informativo"],
    ["V18", "User Agent Fuzzer", "ZAP (10104)", "app-web", "Respuestas distintas según el User-Agent; sin impacto", "Informativo"],
]
tabla("Resumen de vulnerabilidades encontradas",
      ["ID", "Vulnerabilidad", "Fuente", "Activo", "Impacto", "Riesgo"],
      V, [1.0, 4.4, 2.6, 2.4, 3.6, 1.8], tam=8.5, col_riesgo=5,
      nota="Riesgo de ZAP y OpenVAS según la herramienta; riesgo de la revisión de configuración asignado por el grupo.")
tabla("Cantidad de hallazgos por nivel de riesgo",
      ["Riesgo", "Cantidad", "Hallazgos"],
      [["Alto", "2", "V01, V02"], ["Medio", "5", "V03 a V07"], ["Bajo", "7", "V08 a V14"], ["Informativo", "4", "V15 a V18"]],
      [3.0, 2.4, 6.0], col_riesgo=0)

# ================================================================= 8. MITIGACIÓN
H1("8. Actuar: estrategia de mitigación")
P("La estrategia ordena los hallazgos por prioridad —primero los de mayor riesgo y luego los de menor esfuerzo— y "
  "define para cada uno el tratamiento que propone el ciclo: **corregir** la causa, **mitigar** reduciendo su "
  "impacto o probabilidad, o **aceptar** el riesgo cuando es muy bajo o no se puede hacer nada razonable. La "
  "última columna indica cómo se verificaría la corrección.")
M = [
    ["1", "V01", "Corregir", "Mover las contraseñas a un archivo .env (env_file) o a Docker secrets, excluirlo del repositorio y cambiarlas", "docker compose config sin contraseñas visibles; revisión del repositorio"],
    ["2", "V02", "Corregir", "Contraseñas largas y distintas por cuenta (resetPassword); usuario de aplicación con mínimos privilegios", "Intento de ingreso con la clave antigua (debe fallar)"],
    ["3", "V03", "Corregir", "Quitar la publicación del puerto (Oracle solo en soc-net) o limitarla a 127.0.0.1:1521", "docker ps sin 0.0.0.0:1521; escaneo de puertos desde otro equipo"],
    ["4", "V04", "Corregir", "Proxy inverso con TLS (nginx o Traefik); activar SECURE_SSL_REDIRECT, cookies Secure y HSTS en Django", "Nuevo escaneo activo de ZAP sin «HTTP Only Site»"],
    ["5", "V05", "Corregir (aplicado en 1.1)", "Middleware que envía Content-Security-Policy (default-src 'self', frame-ancestors 'none')", "Escaneo de verificación: alerta resuelta (sección 9)"],
    ["6", "V06", "Corregir (aplicado en 1.1)", "WHITENOISE_ALLOW_ALL_ORIGINS = False", "Escaneo de verificación: alerta resuelta (sección 9)"],
    ["7", "V07", "Corregir", "Activar HTTPS en GSA (variable HTTPS de la imagen) y publicar 9392 solo en 127.0.0.1", "Ingreso sin la advertencia «Connection unencrypted»"],
    ["8", "V11", "Corregir (aplicado en 1.1)", "Fijar oracledb==4.0.2 y reconstruir la imagen", "pip en el contenedor muestra 4.0.2; aplicación healthy"],
    ["9", "V12", "Corregir", "Fijar versiones o digest de las imágenes y revisarlas cada mes", "docker compose config / docker images"],
    ["10", "V14", "Mitigar", "Programar la sincronización de feeds fuera de horario (SKIPSYNC=false) o actualizar la imagen cada semana", "Fecha de los feeds en Administration → Feed Status"],
    ["11", "V13", "Mitigar", "Cambiar la clave, restringir el acceso a los registros y pasar la clave por archivo de configuración", "docker logs zap sin la clave"],
    ["12", "V08", "Aceptar", "Riesgo residual documentado: según Django, marcar la cookie CSRF como HttpOnly no aporta protección práctica; la cookie de sesión sí es HttpOnly", "Revisión de Set-Cookie de sessionid (HttpOnly)"],
    ["13", "V09, V10", "Aceptar / mitigar", "Red interna aislada; si se requiere, desactivar net.ipv4.tcp_timestamps y filtrar ICMP timestamp en el firewall del host", "Nueva pasada de OpenVAS"],
    ["14", "V15 a V18", "Aceptar", "Hallazgos informativos sin impacto directo; se mantienen en observación", "Siguiente ciclo de escaneos"],
]
tabla("Estrategia de mitigación priorizada",
      ["Prior.", "ID", "Tratamiento", "Acción propuesta", "Verificación"],
      M, [1.2, 1.4, 2.4, 6.4, 4.4], tam=8.5)

# ================================================================= 9. VERIFICAR Y MEJORAR
H1("9. Verificar y mejorar")
H2("9.1 Aplicación de correcciones y escaneo de verificación")
P("Para cerrar una vuelta del ciclo se aplicaron las tres correcciones de menor esfuerzo sobre la aplicación, sin "
  "tocar el compose de Oracle del docente: un *middleware* que agrega la cabecera Content-Security-Policy (V05), la "
  "configuración `WHITENOISE_ALLOW_ALL_ORIGINS = False` (V06) y el cambio a `oracledb==4.0.2` (V11). Se construyó "
  "la versión `ferreteria-web:1.1` en 27 segundos, reutilizando las capas en caché (Figura 59), y se redesplegó; "
  "las cabeceras confirmaron la CSP y la ausencia de `Access-Control-Allow-Origin` en los estáticos (Figura 60).")
figura("45-verificar-build-app-1-1.png", "Construcción de la versión 1.1 con oracledb 4.0.2.")
figura("46-verificar-app-1-1-cabeceras.png", "Aplicación 1.1 en ejecución: CSP presente y estáticos sin CORS permisivo.")
P("Luego se ejecutó un escaneo de verificación con ZAP en una sesión nueva (Figuras 61 y 62). Las alertas «CSP "
  "Header Not Set» y «Cross-Domain Misconfiguration» desaparecieron. Apareció una alerta nueva de riesgo Medio, "
  "«CSP: style-src unsafe-inline», porque la política permite estilos en línea que usa el panel de administración "
  "de Django. Este resultado muestra por qué el proceso es cíclico: cada corrección se verifica y puede dejar un "
  "punto a mejorar en la siguiente vuelta, en este caso reemplazar los estilos en línea por archivos CSS o usar "
  "*nonces* para quitar `'unsafe-inline'`.")
figura("47-verificar-zap-alertas.png", "Escaneo de verificación: cabecera CSP en la respuesta y alertas restantes.")
figura("48-verificar-zap-reporte3-html.png", "Reporte 3 de ZAP (verificación tras mitigar).", 14)
tabla("Comparación antes y después de las correcciones",
      ["Hallazgo", "App 1.0 (escaneos 1 y 2)", "App 1.1 (verificación)", "Estado"],
      [["V05 Content-Security-Policy no configurada", "Presente (Medio)", "No aparece", "Corregido"],
       ["V06 CORS permisivo en estáticos", "Presente (Medio)", "No aparece", "Corregido"],
       ["V11 oracledb fuera del rango soportado", "26.0.1", "4.0.2", "Corregido"],
       ["CSP: style-src unsafe-inline (10055)", "—", "Presente (Medio)", "Nuevo: siguiente iteración"],
       ["V08 Cookie csrftoken sin HttpOnly", "Presente (Bajo)", "Presente (Bajo)", "Aceptado"]],
      [5.6, 3.6, 3.4, 3.0])
H2("9.2 Métricas")
tabla("Métricas de los escaneos realizados",
      ["Escaneo", "Herramienta", "Hora (Chile)", "Alcance", "Alto", "Medio", "Bajo", "Info."],
      [["Escaneo 1 (pasivo)", "ZAP", "07:14–07:15", "16 URL", "0", "2", "1", "3"],
       ["Escaneo 2 (activo)", "ZAP", "07:17–07:22", "3 444 peticiones", "0", "3", "1", "4"],
       ["Pasada 1", "OpenVAS", "07:23–07:29", "1 host, TCP IANA", "0", "0", "2", "11"],
       ["Pasada 2", "OpenVAS", "07:31–07:36", "1 host, TCP IANA", "0", "0", "2", "11"],
       ["Verificación (app 1.1)", "ZAP", "07:41–07:42", "16 URL", "0", "1", "1", "3"]],
      [3.4, 2.0, 2.4, 3.0, 1.1, 1.2, 1.1, 1.2], tam=9,
      nota="Cantidad de alertas o hallazgos únicos. En OpenVAS, «Info.» corresponde a registros (Log) sin severidad.")
H2("9.3 Mejora continua")
P("Para que el ciclo no dependa de una revisión manual se propone: ejecutar una pasada de OpenVAS cada semana con "
  "un *schedule* de GSA; incorporar el escaneo base de ZAP (`zap-baseline.py`) cada vez que se construya una nueva "
  "versión de la imagen; revisar mensualmente las versiones de las imágenes base y de las dependencias de "
  "`requirements.txt`; mantener actualizados los feeds de OpenVAS; y registrar en cada vuelta las métricas de la "
  "Tabla 11 para observar su tendencia.")

# ================================================================= 10. CONCLUSIONES
H1("10. Conclusiones")
P("El trabajo permitió recorrer una vuelta completa del ciclo de gestión de vulnerabilidades sobre una "
  "infraestructura real de contenedores. Se construyó una aplicación propia, se empaquetó con un Dockerfile, se "
  "desplegó junto a Oracle y a dos escáneres en una red aislada, y se identificaron, evaluaron, priorizaron y "
  "trataron sus vulnerabilidades. Lo más valioso no fue la cantidad de hallazgos, sino haber documentado un "
  "proceso ordenado y repetible: cualquier integrante puede reconstruir la infraestructura a partir de los cuatro "
  "archivos Compose y del Dockerfile, sin depender de comandos recordados de memoria.")
P("Los escáneres confirmaron que la aplicación no presenta vulnerabilidades críticas ni altas: el escaneo activo "
  "de ZAP no logró inyecciones SQL ni XSS, gracias al uso del ORM y del escape automático de plantillas de Django y "
  "a su protección CSRF. Los hallazgos de ZAP se concentraron en la configuración del servidor web (cabeceras, "
  "HTTPS y CORS) y los de OpenVAS fueron de severidad baja, coherentes entre las dos pasadas. Esto refleja que los "
  "marcos de desarrollo modernos protegen bien contra los ataques clásicos, pero no reemplazan una configuración "
  "segura del despliegue.")
P("Los riesgos más altos no fueron detectados por los escáneres, sino por la revisión de configuración: las "
  "credenciales de Oracle en texto plano y la contraseña débil repetida. Esto demuestra que las herramientas "
  "automáticas son necesarias pero no suficientes; ZAP y OpenVAS analizan servicios en ejecución y no leen archivos "
  "de configuración. Un proceso de gestión de vulnerabilidades debe combinar escaneos con revisión manual de la "
  "infraestructura como código.")
P("Los dos incidentes de despliegue —el usuario de Oracle que no se creó y los reinicios en bucle de OpenVAS— "
  "mostraron la importancia de leer los registros de los contenedores y de entender las imágenes que se usan. En "
  "ambos casos la causa estaba documentada en los registros y la solución se aplicó sin improvisar. Del mismo modo, "
  "la verificación de las versiones con la documentación oficial reveló que el driver de Oracle estaba fuera del "
  "rango soportado por Django, un hallazgo que ningún escáner habría reportado.")
P("La fase de verificación confirmó el carácter cíclico del proceso: tras aplicar tres correcciones y volver a "
  "escanear, dos alertas desaparecieron y apareció una nueva derivada de la propia corrección. Que hoy una "
  "aplicación no tenga vulnerabilidades no garantiza que no las tenga mañana, cuando se publiquen nuevas fallas en "
  "Python, Django u Oracle. Por eso se propuso automatizar escaneos periódicos y revisar las versiones en cada "
  "vuelta.")
P("Finalmente, Docker demostró ser una base adecuada para este tipo de trabajo: permitió levantar en minutos "
  "herramientas complejas como OpenVAS o ZAP sin instalarlas en el equipo, aislar los servicios en una red propia y "
  "repetir el despliegue de forma idéntica. Como propuesta de profundización queda implementar HTTPS con un proxy "
  "inverso, mover los secretos a Docker secrets y desplegar la solución en un proveedor en la nube, como mencionó "
  "el docente.")

# ================================================================= REFERENCIAS
H1("Referencias")
REFS = [
    "CrowdStrike. (s.f.). *¿Qué es el ciclo de vida de la gestión de vulnerabilidades?* Recuperado el 27 de octubre de 2026, de https://www.crowdstrike.com/es-es/cybersecurity-101/exposure-management/vulnerability-management-lifecycle/",
    "Django Software Foundation. (s.f.-a). *Databases* (Django 5.2). Recuperado el 27 de octubre de 2026, de https://docs.djangoproject.com/en/5.2/ref/databases/",
    "Django Software Foundation. (s.f.-b). *Settings: CSRF_COOKIE_HTTPONLY* (Django 5.2). Recuperado el 27 de octubre de 2026, de https://docs.djangoproject.com/en/5.2/ref/settings/",
    "Docker Inc. (s.f.-a). *Docker overview*. Recuperado el 27 de octubre de 2026, de https://docs.docker.com/get-started/docker-overview/",
    "Docker Inc. (s.f.-b). *Networking overview*. Recuperado el 27 de octubre de 2026, de https://docs.docker.com/engine/network/",
    "Docker Inc. (s.f.-c). *Volumes*. Recuperado el 27 de octubre de 2026, de https://docs.docker.com/engine/storage/volumes/",
    "Docker Inc. (s.f.-d). *Dockerfile reference*. Recuperado el 27 de octubre de 2026, de https://docs.docker.com/reference/dockerfile/",
    "Docker Inc. (s.f.-e). *Docker Compose*. Recuperado el 27 de octubre de 2026, de https://docs.docker.com/compose/",
    "Greenbone AG. (s.f.). *Greenbone Community Documentation*. Recuperado el 27 de octubre de 2026, de https://greenbone.github.io/docs/latest/",
    "Immauss. (s.f.). *OpenVAS docker container*. Recuperado el 27 de octubre de 2026, de https://immauss.github.io/openvas/",
    "Oracle. (s.f.). *Initializing python-oracledb*. Recuperado el 27 de octubre de 2026, de https://python-oracledb.readthedocs.io/en/latest/user_guide/initialization.html",
    "Souppaya, M. y Scarfone, K. (2022). *Guide to enterprise patch management planning: Preventive maintenance for technology* (NIST Special Publication 800-40 Rev. 4). National Institute of Standards and Technology. https://doi.org/10.6028/NIST.SP.800-40r4",
    "ZAP. (s.f.-a). *ZAP – Docker user guide*. Recuperado el 27 de octubre de 2026, de https://www.zaproxy.org/docs/docker/about/",
    "ZAP. (s.f.-b). *ZAP – Webswing*. Recuperado el 27 de octubre de 2026, de https://www.zaproxy.org/docs/docker/webswing/",
    "ZAP. (s.f.-c). *Content Security Policy (CSP) Header Not Set* (alerta 10038). Recuperado el 27 de octubre de 2026, de https://www.zaproxy.org/docs/alerts/10038/",
]
for ref in REFS:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)        # sangría francesa (APA)
    runs(p, ref)

# ================================================================= ANEXOS
H1("Anexos")
H2("Anexo A. Archivos docker-compose.yml")
for t, ruta in (("A.1 Oracle (01-oracle/docker-compose.yml, entregado por el docente)", "01-oracle/docker-compose.yml"),
                ("A.2 Aplicación web (02-app-django/docker-compose.yml, versión final con la imagen 1.1)", "02-app-django/docker-compose.yml"),
                ("A.3 ZAP (03-zap/docker-compose.yml)", "03-zap/docker-compose.yml"),
                ("A.4 OpenVAS (04-openvas/docker-compose.yml, versión final con SKIPSYNC)", "04-openvas/docker-compose.yml")):
    H3(t)
    codigo(leer(ruta))
H2("Anexo B. Dockerfile de la aplicación web")
codigo(leer("02-app-django/Dockerfile"))
H2("Anexo C. Comandos de la demostración")
codigo("""# Contenedores en ejecución
docker ps --format "table {{.Names}}\\t{{.Image}}\\t{{.Status}}\\t{{.Ports}}"

# Contenedores de soc-net con su IP (comando del docente; en PowerShell escapar las comillas internas)
docker network inspect soc-net -f '{{range .Containers}}{{.Name}} - {{.IPv4Address}}{{"\\n"}}{{end}}'
docker network inspect soc-net -f '{{range .Containers}}{{.Name}} - {{.IPv4Address}}{{\\"\\n\\"}}{{end}}'

# Interfaces web
http://localhost:8088        aplicación web (Django)
http://localhost:8080/zap/   ZAP
http://localhost:9392        OpenVAS (GSA)""")

salida.parent.mkdir(parents=True, exist_ok=True)
doc.save(salida)
print("Informe generado:", salida)
